"""
Word (.docx) faylidan savollarni parse qilish.
word_import.py va QuestionImportView ikkalasi ham shu modulni ishlatadi.
"""
from docx import Document


# ─── yordamchi ────────────────────────────────────────────────────────────────

def _cell(table, row, col):
    try:
        return table.cell(row, col).text.strip()
    except IndexError:
        return ""

def _key(table, row):
    return _cell(table, row, 0).upper().strip()

def _val(table, row):
    return _cell(table, row, 1)


# ─── parsers ──────────────────────────────────────────────────────────────────

def _parse_variantli(table):
    savol, options, correct, points, izoh = "", [], "", 1, ""
    for i in range(1, len(table.rows)):
        k, v = _key(table, i), _val(table, i)
        if k == "SAVOL":
            savol = v
        elif k in ("A","B","C","D","E","F","A*","B*","C*","D*","E*","F*"):
            if v:
                options.append(v)
                if k.endswith("*"):
                    correct = v
        elif k == "IZOH":
            izoh = v
        elif k == "BALL":
            try:
                points = int(v)
            except ValueError:
                pass

    if not savol:
        return None, "SAVOL bo'sh"
    if len(options) < 2:
        return None, "Kamida 2 ta variant kerak"
    if not correct:
        return None, "To'g'ri javob ko'rsatilmagan (masalan: B*)"

    q = {"question_text": savol,
         "question_data": {"options": options},
         "correct_answer": correct,
         "points": points}
    if izoh:
        q["explanation"] = izoh
    return q, None


def _parse_ha_yoq(table):
    savol, javob, points, izoh = "", "", 1, ""
    for i in range(1, len(table.rows)):
        k, v = _key(table, i), _val(table, i)
        if k == "SAVOL":
            savol = v
        elif k == "JAVOB":
            raw = v.upper().strip()
            if raw in ("HA", "TO'G'RI", "TRUE", "1", "+"):
                javob = "true"
            elif raw in ("YO'Q", "YOQ", "NOTO'G'RI", "FALSE", "0", "-"):
                javob = "false"
        elif k == "IZOH":
            izoh = v
        elif k == "BALL":
            try:
                points = int(v)
            except ValueError:
                pass

    if not savol:
        return None, "SAVOL bo'sh"
    if not javob:
        return None, "JAVOB ko'rsatilmagan (HA yoki YO'Q)"

    q = {"question_text": savol,
         "question_data": {"statement": savol},
         "correct_answer": javob,
         "points": points}
    if izoh:
        q["explanation"] = izoh
    return q, None


def _parse_bosh(table):
    matn, javob, points, izoh = "", "", 1, ""
    for i in range(1, len(table.rows)):
        k, v = _key(table, i), _val(table, i)
        if k == "MATN":
            matn = v
        elif k in ("JAVOB", "JAVOB1"):
            javob = v
        elif k == "IZOH":
            izoh = v
        elif k == "BALL":
            try:
                points = int(v)
            except ValueError:
                pass

    if not matn:
        return None, "MATN bo'sh"
    if not javob:
        return None, "JAVOB ko'rsatilmagan"
    if "___" not in matn:
        return None, "MATNda ___ (bo'sh joy belgisi) yo'q"

    q = {"question_text": matn,
         "question_data": {"text": matn, "blanks": matn.count("___")},
         "correct_answer": javob,
         "points": points}
    if izoh:
        q["explanation"] = izoh
    return q, None


def _parse_moslashtirish(table):
    RESERVED = {"MOSLASHTIRISH", "MATCHING", "SAVOL", "BALL", "IZOH", ""}
    savol = "Quyidagilarni mos keltiring:"
    pairs, correct_map, points, izoh = [], {}, 1, ""

    for i in range(1, len(table.rows)):
        k             = _key(table, i)
        v             = _val(table, i)
        original_left = _cell(table, i, 0).strip()

        if k == "SAVOL":
            savol = v or savol
        elif k == "IZOH":
            izoh = v
        elif k == "BALL":
            try:
                points = int(v)
            except ValueError:
                pass
        elif k not in RESERVED and original_left and v:
            pairs.append({"left": original_left, "right": v})
            correct_map[original_left] = v

    if len(pairs) < 2:
        return None, "Kamida 2 ta juftlik kerak"

    q = {"question_text": savol,
         "question_data": {"pairs": pairs},
         "correct_answer": correct_map,
         "points": points}
    if izoh:
        q["explanation"] = izoh
    return q, None


# ─── tur xaritasi ─────────────────────────────────────────────────────────────

_PARSERS = {
    "VARIANTLI":          _parse_variantli,
    "KO'P TANLOV":        _parse_variantli,
    "MULTIPLE CHOICE":    _parse_variantli,
    "HA_YOQ":             _parse_ha_yoq,
    "HA/YO'Q":            _parse_ha_yoq,
    "HA YO'Q":            _parse_ha_yoq,
    "TRUE FALSE":         _parse_ha_yoq,
    "TO'G'RI/NOTO'G'RI": _parse_ha_yoq,
    "BO'SH":              _parse_bosh,
    "BO'SH TO'LDIRISH":  _parse_bosh,
    "FILLING":            _parse_bosh,
    "FILL BLANK":         _parse_bosh,
    "MOSLASHTIRISH":      _parse_moslashtirish,
    "MATCHING":           _parse_moslashtirish,
}


# ─── asosiy funksiya ──────────────────────────────────────────────────────────

def parse_docx(file_obj):
    """
    file_obj — ochiq fayl yoki BytesIO obyekti.

    Qaytaradi:
        questions : list[dict]   — muvaffaqiyatli parse qilingan savollar
        errors    : list[dict]   — xato ma'lumotlari { jadval, tur, xato }
    """
    doc = Document(file_obj)
    questions, errors = [], []

    for idx, table in enumerate(doc.tables, 1):
        if len(table.rows) < 2:
            continue

        type_raw = _cell(table, 0, 0).strip().upper()
        parser   = _PARSERS.get(type_raw)

        if parser is None:
            errors.append({
                "jadval": idx,
                "xato": f"Noma'lum tur: '{type_raw}'"
            })
            continue

        q, err = parser(table)
        if err:
            errors.append({"jadval": idx, "tur": type_raw, "xato": err})
            continue

        q["order_index"] = len(questions) + 1
        questions.append(q)

    return questions, errors
